from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from models import ChartOfAccount, Employee, Subcontractor, Supplier, db
from services.accounting import (
    ENTITY_KIND_OPTIONS, EXPENSE_CLASS_OPTIONS, as_float, sync_journal_related_accounts,
)
from services.immutability import account_has_posted_moves

SHEET_ACCOUNTS = "الحسابات"
SHEET_SUPPLIERS = "الموردون"
SHEET_SUBS = "مقاولو الباطن"
SHEET_EMPLOYEES = "الموظفون"

ACCOUNT_HEADERS = ["الكود", "الاسم", "الفئة", "الرصيد الافتتاحي", "تبويب المصروف"]
SUPPLIER_HEADERS = ["الاسم", "الكود", "التصنيف", "الاتصال", "ملاحظات"]
SUB_HEADERS = ["الاسم", "الكود", "التصنيف", "الاتصال", "ملاحظات"]
EMPLOYEE_HEADERS = ["الاسم", "الميزة", "المرتب الأساسي", "الاتصال", "ملاحظات"]

HEADER_FILL = PatternFill("solid", fgColor="1E73E8")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _cell_text(value):
    if value is None:
        return ""
    return str(value).strip()


def _sheet(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.rightToLeft = True
    sheet.append(headers)
    for col in range(1, len(headers) + 1):
        cell = sheet.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    data = rows or []
    for row in data:
        sheet.append(list(row))
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 22
    return sheet


def build_import_workbook(include_existing=False):
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    account_rows = []
    supplier_rows = []
    sub_rows = []
    employee_rows = []
    if include_existing:
        for account in ChartOfAccount.query.order_by(ChartOfAccount.category, ChartOfAccount.code).all():
            account_rows.append([
                account.code,
                account.name,
                account.category,
                account.opening_balance or 0,
                account.expense_class or "",
            ])
        for supplier in Supplier.query.order_by(Supplier.name).all():
            supplier_rows.append([
                supplier.name,
                supplier.code or "",
                supplier.entity_kind or "",
                supplier.contact_info or "",
                supplier.notes or "",
            ])
        for sub in Subcontractor.query.order_by(Subcontractor.name).all():
            sub_rows.append([
                sub.name,
                sub.code or "",
                sub.entity_kind or "",
                sub.contact_info or "",
                sub.notes or "",
            ])
        for employee in Employee.query.order_by(Employee.name).all():
            employee_rows.append([
                employee.name,
                employee.tag or "موظف",
                employee.basic_salary or 0,
                employee.contact_info or "",
                employee.notes or "",
            ])

    _sheet(workbook, SHEET_ACCOUNTS, ACCOUNT_HEADERS, account_rows)
    _sheet(workbook, SHEET_SUPPLIERS, SUPPLIER_HEADERS, supplier_rows)
    _sheet(workbook, SHEET_SUBS, SUB_HEADERS, sub_rows)
    _sheet(workbook, SHEET_EMPLOYEES, EMPLOYEE_HEADERS, employee_rows)
    return workbook


def workbook_bytes(workbook):
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _header_map(sheet):
    mapping = {}
    for index, cell in enumerate(sheet[1], start=1):
        key = _cell_text(cell.value)
        if key:
            mapping[key] = index
    return mapping


def _row_values(sheet, row_index, headers, mapping):
    values = {}
    for header in headers:
        col = mapping.get(header)
        values[header] = sheet.cell(row_index, col).value if col else None
    return values


def _is_empty_row(values):
    return not any(_cell_text(value) for value in values.values())


def import_workbook_file(file_storage):
    workbook = load_workbook(file_storage, data_only=True)
    summary = {
        "accounts_added": 0,
        "accounts_updated": 0,
        "suppliers_added": 0,
        "subs_added": 0,
        "employees_added": 0,
        "skipped": [],
    }

    if SHEET_ACCOUNTS in workbook.sheetnames:
        _import_accounts(workbook[SHEET_ACCOUNTS], summary)
    if SHEET_SUPPLIERS in workbook.sheetnames:
        _import_suppliers(workbook[SHEET_SUPPLIERS], summary)
    if SHEET_SUBS in workbook.sheetnames:
        _import_subs(workbook[SHEET_SUBS], summary)
    if SHEET_EMPLOYEES in workbook.sheetnames:
        _import_employees(workbook[SHEET_EMPLOYEES], summary)

    db.session.commit()
    sync_journal_related_accounts()
    return summary


def _import_accounts(sheet, summary):
    mapping = _header_map(sheet)
    if "الكود" not in mapping or "الاسم" not in mapping:
        summary["skipped"].append("ورقة الحسابات: العناوين غير مطابقة للنموذج")
        return
    for row_index in range(2, sheet.max_row + 1):
        values = _row_values(sheet, row_index, ACCOUNT_HEADERS, mapping)
        if _is_empty_row(values):
            continue
        code = _cell_text(values["الكود"])
        name = _cell_text(values["الاسم"])
        category = _cell_text(values["الفئة"]) or "المصروفات"
        if not code or not name:
            summary["skipped"].append(f"حساب صف {row_index}: ناقص الكود أو الاسم")
            continue
        opening = as_float(values["الرصيد الافتتاحي"])
        expense_class = _cell_text(values["تبويب المصروف"])
        if expense_class and expense_class not in EXPENSE_CLASS_OPTIONS:
            expense_class = None
        existing = ChartOfAccount.query.filter_by(code=code).first()
        if existing:
            existing.name = name
            existing.category = category
            if expense_class:
                existing.expense_class = expense_class
            if not account_has_posted_moves(existing.id):
                existing.opening_balance = opening
            summary["accounts_updated"] += 1
            continue
        db.session.add(ChartOfAccount(
            code=code,
            name=name,
            category=category,
            opening_balance=opening,
            expense_class=expense_class,
        ))
        summary["accounts_added"] += 1


def _find_party(model, name, code):
    if code:
        found = model.query.filter_by(code=code).first()
        if found:
            return found
    return model.query.filter_by(name=name).first()


def _import_suppliers(sheet, summary):
    mapping = _header_map(sheet)
    if "الاسم" not in mapping:
        summary["skipped"].append("ورقة الموردون: عنوان الاسم غير موجود")
        return
    for row_index in range(2, sheet.max_row + 1):
        values = _row_values(sheet, row_index, SUPPLIER_HEADERS, mapping)
        if _is_empty_row(values):
            continue
        name = _cell_text(values["الاسم"])
        if not name:
            continue
        code = _cell_text(values["الكود"]) or None
        kind = _cell_text(values["التصنيف"]) or ENTITY_KIND_OPTIONS[1]
        if kind not in ENTITY_KIND_OPTIONS:
            kind = ENTITY_KIND_OPTIONS[1]
        existing = _find_party(Supplier, name, code)
        if existing:
            existing.name = name
            if code:
                existing.code = code
            existing.entity_kind = kind
            existing.contact_info = _cell_text(values["الاتصال"]) or None
            existing.notes = _cell_text(values["ملاحظات"]) or None
            continue
        supplier = Supplier(
            name=name,
            code=code,
            entity_kind=kind,
            contact_info=_cell_text(values["الاتصال"]) or None,
            notes=_cell_text(values["ملاحظات"]) or None,
        )
        db.session.add(supplier)
        db.session.flush()
        if not supplier.code:
            supplier.code = f"SUP-{supplier.id:04d}"
        summary["suppliers_added"] += 1


def _import_subs(sheet, summary):
    mapping = _header_map(sheet)
    if "الاسم" not in mapping:
        summary["skipped"].append("ورقة مقاولو الباطن: عنوان الاسم غير موجود")
        return
    for row_index in range(2, sheet.max_row + 1):
        values = _row_values(sheet, row_index, SUB_HEADERS, mapping)
        if _is_empty_row(values):
            continue
        name = _cell_text(values["الاسم"])
        if not name:
            continue
        code = _cell_text(values["الكود"]) or None
        kind = _cell_text(values["التصنيف"]) or ENTITY_KIND_OPTIONS[0]
        if kind not in ENTITY_KIND_OPTIONS:
            kind = ENTITY_KIND_OPTIONS[0]
        existing = _find_party(Subcontractor, name, code)
        if existing:
            existing.name = name
            if code:
                existing.code = code
            existing.entity_kind = kind
            existing.contact_info = _cell_text(values["الاتصال"]) or None
            existing.notes = _cell_text(values["ملاحظات"]) or None
            continue
        sub = Subcontractor(
            name=name,
            code=code,
            entity_kind=kind,
            contact_info=_cell_text(values["الاتصال"]) or None,
            notes=_cell_text(values["ملاحظات"]) or None,
        )
        db.session.add(sub)
        db.session.flush()
        if not sub.code:
            sub.code = f"SUB-{sub.id:04d}"
        summary["subs_added"] += 1


def _import_employees(sheet, summary):
    mapping = _header_map(sheet)
    if "الاسم" not in mapping:
        summary["skipped"].append("ورقة الموظفون: عنوان الاسم غير موجود")
        return
    for row_index in range(2, sheet.max_row + 1):
        values = _row_values(sheet, row_index, EMPLOYEE_HEADERS, mapping)
        if _is_empty_row(values):
            continue
        name = _cell_text(values["الاسم"])
        if not name:
            continue
        tag = _cell_text(values["الميزة"]) or "موظف"
        existing = Employee.query.filter_by(name=name).first()
        if existing:
            existing.tag = tag
            existing.basic_salary = as_float(values["المرتب الأساسي"])
            existing.contact_info = _cell_text(values["الاتصال"]) or None
            existing.notes = _cell_text(values["ملاحظات"]) or None
            continue
        employee = Employee(
            name=name,
            tag=tag,
            basic_salary=as_float(values["المرتب الأساسي"]),
            contact_info=_cell_text(values["الاتصال"]) or None,
            notes=_cell_text(values["ملاحظات"]) or None,
        )
        db.session.add(employee)
        db.session.flush()
        if not employee.code:
            employee.code = f"EMP-{employee.id:04d}"
        summary["employees_added"] += 1
