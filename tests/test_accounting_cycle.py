import os
import tempfile

import pytest

_DB_FD, _DB_PATH = tempfile.mkstemp(suffix=".db")
os.close(_DB_FD)
os.environ["SECRET_KEY"] = "test-secret-key"
os.environ["DATABASE_URL"] = "sqlite:///" + _DB_PATH.replace("\\", "/")
os.environ.pop("RENDER", None)
os.environ["FLASK_ENV"] = "testing"

from app import app as flask_app
from models import (
    ROLE_ADMIN, ROLE_DATA_ENTRY, AccountingPeriod, ActivityLog, ChartOfAccount, Estimation, Equipment, JournalEntry,
    ProgressPayment, ProgressPaymentItem, Project, Subcontractor, Supplier, User, db, PurchaseOrder,
    InventoryTransaction, ClientReceipt, SubcontractorPayment, SupplierPayment, Employee,
    PayrollSlip, EmployeeAttendance, EmployeeSalaryPayment, SubcontractorRetentionRelease, SalesTrip,
)
from services.accounting import (
    build_account_balances, build_project_cost_breakdown, build_subcontractor_statement,
    generate_progress_payment_number, get_account_by_code, get_or_create_employee_account,
    get_subcontractor_outstanding_advances, get_subcontractor_outstanding_retention,
    previous_executed_quantity, progress_net_remaining, sync_equipment_journals,
    sync_journal_related_accounts,
)


@pytest.fixture(autouse=True)
def _db():
    flask_app.config["TESTING"] = True
    flask_app.config["WTF_CSRF_ENABLED"] = False
    with flask_app.app_context():
        db.drop_all()
        db.create_all()
        sync_journal_related_accounts()
        admin = User(username="admin", full_name="مدير النظام", role=ROLE_ADMIN, is_active=True)
        admin.set_password("secret12")
        db.session.add(admin)
        db.session.commit()
        yield
        db.session.remove()


@pytest.fixture()
def client():
    test_client = flask_app.test_client()
    with flask_app.app_context():
        admin = User.query.filter_by(username="admin").first()
        admin_id = admin.id
    with test_client.session_transaction() as sess:
        sess["user_id"] = admin_id
        sess["admin_amounts_visible"] = True
    return test_client


def _ids():
    treasury = ChartOfAccount.query.filter_by(code="TRS-MAIN").first()
    expense = ChartOfAccount.query.filter_by(code="EXP-ELC").first() or ChartOfAccount.query.filter_by(category="المصروفات").first()
    return treasury.id, expense.id


def test_client_sales_exclude_subcontractor_ipc(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-T1", project_name="برج", client_name="عميل أ",
            contract_type="مقاولات عامة", contract_value=1000000,
        )
        sub = Subcontractor(name="مقاول اختبار")
        db.session.add_all([project, sub])
        db.session.commit()
        pid, sid = project.id, sub.id

    client.post("/progress_payments", data={
        "project_id": str(pid),
        "subcontractor_id": str(sid),
        "date": "2026-08-01",
        "retention_percentage": "0",
        "tax_percentage": "0",
        "description": ["خرسانة"],
        "unit": ["متر مكعب"],
        "quantity": ["10"],
        "unit_price": ["1000"],
    }, follow_redirects=True)

    client.post("/client_progress_payments", data={
        "project_id": str(pid),
        "date": "2026-08-02",
        "retention_percentage": "0",
        "tax_percentage": "0",
        "description": ["أعمال للعميل"],
        "unit": ["مقطوعية"],
        "quantity": ["1"],
        "unit_price": ["50000"],
    }, follow_redirects=True)

    home = client.get("/")
    body = home.get_data(as_text=True)
    assert "50000" in body.replace(",", "").replace(".", "") or "50.000" in body
    with flask_app.app_context():
        assert ProgressPayment.query.filter_by(subcontractor_id=sid).count() == 1
        assert ProgressPayment.query.filter(ProgressPayment.subcontractor_id.is_(None)).count() == 1


def test_client_progress_payment_computes_line_total_and_net(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-CLI-TOT", project_name="نقل رمل", client_name="عميل إجمالي",
            contract_type="مقاولات عامة", contract_value=1,
        )
        db.session.add(project)
        db.session.commit()
        pid = project.id

    page = client.get("/client_progress_payments")
    body = page.get_data(as_text=True)
    assert "الإجمالي المستحق" in body
    assert "line-total" in body
    assert "ipc-line" in body
    assert "صافي المستحق على العميل (كشف الحساب)" in body

    client.post("/client_progress_payments", data={
        "project_id": str(pid),
        "date": "2026-08-04",
        "retention_percentage": "10",
        "tax_percentage": "0",
        "description": ["نقلة رمل", "يومية سائق"],
        "unit": ["نقلة", "يومية"],
        "quantity": ["5", "2"],
        "unit_price": ["1000", "500"],
    }, follow_redirects=True)

    with flask_app.app_context():
        payment = ProgressPayment.query.filter(ProgressPayment.subcontractor_id.is_(None)).one()
        items = ProgressPaymentItem.query.filter_by(progress_payment_id=payment.id).all()
        assert sorted(round(item.value, 2) for item in items) == [1000.0, 5000.0]
        assert round(payment.total_value, 2) == 6000.0
        assert round(payment.discount_insurance, 2) == 600.0
        assert round(payment.net_value, 2) == 5400.0


def test_estimation_does_not_post_revenue(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-T2", project_name="فيلا", client_name="عميل ب",
            contract_type="مقاولات عامة", contract_value=200000,
        )
        db.session.add(project)
        db.session.commit()
        pid = project.id

    client.post("/estimations", data={
        "client_name": "عميل ب",
        "project_id": str(pid),
        "date": "2026-08-01",
        "discount_percentage": "0",
        "admin_percentage": "0",
        "admin_mode": "إضافة",
        "item_description": ["بند"],
        "item_unit": ["مقطوعية"],
        "item_quantity": ["1"],
        "item_unit_price": ["80000"],
        "item_discount_percentage": ["0"],
    }, follow_redirects=True)

    with flask_app.app_context():
        estimation = Estimation.query.first()
        est_id = estimation.id

    client.post(f"/estimations/{est_id}/status", data={"status": "معتمدة"}, follow_redirects=True)
    with flask_app.app_context():
        autos = JournalEntry.query.filter(JournalEntry.description.like("%EST-AUTO:%")).all()
        assert autos == []
        revenue = ChartOfAccount.query.filter_by(code="REV-WRK").first()
        balances = build_account_balances()
        assert round(balances.get(revenue.id, 0), 2) == 0.0


def test_purchase_order_does_not_double_inventory(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-T3", project_name="مخزن", client_name="عميل ج",
            contract_type="مقاولات عامة", contract_value=1,
        )
        supplier = Supplier(name="مورد حديد")
        db.session.add_all([project, supplier])
        db.session.commit()
        pid, sid = project.id, supplier.id

    client.post("/purchase_orders", data={
        "project_id": str(pid),
        "supplier_id": str(sid),
        "item_name": "حديد",
        "warehouse_name": "رئيسي",
        "quantity": "10",
        "unit_price": "100",
        "discount": "0",
        "date": "2026-08-01",
        "status": "مفتوح",
    }, follow_redirects=True)

    with flask_app.app_context():
        order = PurchaseOrder.query.first()
        assert order is not None
        tx = InventoryTransaction.query.filter(InventoryTransaction.notes.like(f"%PO-AUTO:{order.id}%")).first()
        po_journals = JournalEntry.query.filter(JournalEntry.description.like(f"%PO-JRN-AUTO:{order.id}%")).all()
        assert tx is None
        assert po_journals == []
        oid = order.id

    client.post(f"/purchase_orders/{oid}/update", data={
        "project_id": str(pid),
        "supplier_id": str(sid),
        "item_name": "حديد",
        "warehouse_name": "رئيسي",
        "quantity": "10",
        "unit_price": "100",
        "discount": "0",
        "date": "2026-08-01",
        "status": "مغلق",
    }, follow_redirects=True)

    with flask_app.app_context():
        tx = InventoryTransaction.query.filter(InventoryTransaction.notes.like(f"%PO-AUTO:{oid}%")).first()
        assert tx is not None
        assert round(tx.quantity, 2) == 10
        inv_journals = JournalEntry.query.filter(JournalEntry.description.like(f"%INV-AUTO:{tx.id}%")).all()
        po_journals = JournalEntry.query.filter(JournalEntry.description.like(f"%PO-JRN-AUTO:{oid}%")).all()
        assert po_journals and round(po_journals[0].amount, 2) == 1000.0
        assert inv_journals == []


def test_inventory_transfer_and_stock_guard(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-T3B", project_name="مخزن ب", client_name="عميل",
            contract_type="مقاولات عامة", contract_value=1,
        )
        supplier = Supplier(name="مورد أسمنت")
        db.session.add_all([project, supplier])
        db.session.commit()
        pid, sid = project.id, supplier.id

    client.post("/purchase_orders", data={
        "project_id": str(pid),
        "supplier_id": str(sid),
        "item_name": ["أسمنت", "رمل"],
        "quantity": ["20", "5"],
        "unit_price": ["50", "10"],
        "discount": ["0", "0"],
        "warehouse_name": "رئيسي",
        "date": "2026-08-01",
        "status": "مغلق",
    }, follow_redirects=True)

    over_issue = client.post("/inventory", data={
        "project_id": str(pid),
        "material_name": "أسمنت",
        "warehouse_name": "رئيسي",
        "quantity": "21",
        "unit_cost": "50",
        "transaction_type": "سحب",
        "date": "2026-08-02",
    }, follow_redirects=True)
    assert "غير متاحة" in over_issue.get_data(as_text=True)

    client.post("/inventory", data={
        "project_id": str(pid),
        "material_name": "أسمنت",
        "warehouse_name": "رئيسي",
        "destination_warehouse": "موقع",
        "quantity": "8",
        "transaction_type": "تحويل",
        "date": "2026-08-03",
    }, follow_redirects=True)

    with flask_app.app_context():
        from services.accounting import available_stock_qty
        assert round(available_stock_qty(pid, "رئيسي", "أسمنت"), 2) == 12
        assert round(available_stock_qty(pid, "موقع", "أسمنت"), 2) == 8
        order = PurchaseOrder.query.first()
        oid = order.id

    deleted = client.post(f"/documents/purchase_order/{oid}/delete", follow_redirects=True)
    assert "لا يمكن تخفيض استلام" in deleted.get_data(as_text=True)
    with flask_app.app_context():
        assert PurchaseOrder.query.get(oid) is not None
        assert InventoryTransaction.query.filter_by(transaction_type="تحويل").count() == 1


def test_delete_unused_purchase_order_clears_stock(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-T3C", project_name="مخزن ج", client_name="عميل",
            contract_type="مقاولات عامة", contract_value=1,
        )
        supplier = Supplier(name="مورد رمل")
        db.session.add_all([project, supplier])
        db.session.commit()
        pid, sid = project.id, supplier.id

    client.post("/purchase_orders", data={
        "project_id": str(pid),
        "supplier_id": str(sid),
        "item_name": "رمل",
        "warehouse_name": "رئيسي",
        "quantity": "4",
        "unit_price": "20",
        "discount": "0",
        "date": "2026-08-01",
        "status": "مغلق",
    }, follow_redirects=True)

    with flask_app.app_context():
        order = PurchaseOrder.query.first()
        oid = order.id
        assert InventoryTransaction.query.filter(InventoryTransaction.notes.like(f"%PO-AUTO:{oid}%")).first() is not None

    deleted = client.post(f"/documents/purchase_order/{oid}/delete", follow_redirects=True)
    assert deleted.status_code == 200
    with flask_app.app_context():
        assert PurchaseOrder.query.get(oid) is None
        assert InventoryTransaction.query.filter(InventoryTransaction.notes.like(f"%PO-AUTO:{oid}%")).all() == []
        assert JournalEntry.query.filter(JournalEntry.description.like(f"%PO-JRN-AUTO:{oid}%")).all() == []


def test_project_cost_matches_profitability_source(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-T4", project_name="طريق", client_name="عميل د",
            contract_type="مقاولات عامة", contract_value=1, admin_percentage=10,
        )
        sub = Subcontractor(name="باطن")
        db.session.add_all([project, sub])
        db.session.commit()
        from models import ProgressPayment, ProgressPaymentItem
        payment = ProgressPayment(project_id=project.id, subcontractor_id=sub.id, date="2026-08-01", total_value=0)
        db.session.add(payment)
        db.session.commit()
        db.session.add(ProgressPaymentItem(
            progress_payment_id=payment.id, description="أعمال", unit="م3",
            quantity=2, unit_price=1000, value=2000,
        ))
        from services.accounting import recalculate_progress_payment, sync_progress_payment_journals
        recalculate_progress_payment(payment)
        db.session.commit()
        breakdown = build_project_cost_breakdown(project)
        assert breakdown["subcontractor_cost"] == 2000.0
        assert breakdown["admin_allocation"] == 200.0
        assert breakdown["total_cost"] == 2200.0


def test_client_receipt_posts_to_receivable_and_treasury(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-T5", project_name="تحصيل", client_name="عميل هـ",
            contract_type="مقاولات عامة", contract_value=1,
        )
        db.session.add(project)
        db.session.commit()
        pid = project.id
        treasury_id, _ = _ids()

    client.post("/client_receipts", data={
        "client_name": "عميل هـ",
        "project_id": str(pid),
        "amount": "1500",
        "payment_method": "نقدي",
        "treasury_account_id": str(treasury_id),
        "date": "2026-08-10",
    }, follow_redirects=True)

    with flask_app.app_context():
        receipt = ClientReceipt.query.first()
        assert receipt is not None
        journals = JournalEntry.query.filter(JournalEntry.description.like(f"%REC-AUTO:{receipt.id}%")).all()
        assert journals and round(journals[0].amount, 2) == 1500.0
        client_account = ChartOfAccount.query.filter_by(name="عميل - عميل هـ").first()
        treasury = ChartOfAccount.query.filter_by(code="TRS-MAIN").first()
        assert journals[0].debit_account_id == treasury.id
        assert journals[0].credit_account_id == client_account.id


def test_closed_period_blocks_journal_post(client):
    with flask_app.app_context():
        period = AccountingPeriod(name="أغسطس", from_date="2026-08-01", to_date="2026-08-31", status="مغلقة")
        db.session.add(period)
        db.session.commit()
        treasury_id, expense_id = _ids()

    response = client.post("/journal", data={
        "date": "2026-08-15",
        "entry_action": "post",
        "line_description": ["قيد تجريبي"],
        "line_debit_account_id": [str(expense_id)],
        "line_credit_account_id": [str(treasury_id)],
        "line_amount": ["100"],
    }, follow_redirects=True)
    body = response.get_data(as_text=True)
    assert "مغلقة" in body
    with flask_app.app_context():
        assert JournalEntry.query.count() == 0


def test_data_entry_cannot_open_accounts(client):
    with flask_app.app_context():
        clerk = User(username="clerk", full_name="مدخل بيانات", role=ROLE_DATA_ENTRY, is_active=True)
        clerk.set_password("secret12")
        db.session.add(clerk)
        db.session.commit()
        clerk_id = clerk.id

    probe = flask_app.test_client()
    with probe.session_transaction() as sess:
        sess["user_id"] = clerk_id
    response = probe.get("/accounts", follow_redirects=True)
    assert "ليست لديك صلاحية" in response.get_data(as_text=True)


def test_posted_documents_delete_clears_both_accounts(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-T6", project_name="حذف", client_name="عميل و",
            contract_type="مقاولات عامة", contract_value=1,
        )
        db.session.add(project)
        db.session.commit()
        pid = project.id

    client.post("/client_progress_payments", data={
        "project_id": str(pid),
        "date": "2026-08-03",
        "retention_percentage": "0",
        "tax_percentage": "0",
        "description": ["بند"],
        "unit": ["مقطوعية"],
        "quantity": ["1"],
        "unit_price": ["9000"],
    }, follow_redirects=True)

    with flask_app.app_context():
        payment = ProgressPayment.query.filter(ProgressPayment.subcontractor_id.is_(None)).first()
        payment_id = payment.id
        journals = JournalEntry.query.filter(JournalEntry.description.like(f"%PP-AUTO:{payment_id}%")).all()
        assert len(journals) == 1
        debit_id = journals[0].debit_account_id
        credit_id = journals[0].credit_account_id

    response = client.post(f"/documents/progress_payment/{payment_id}/delete", follow_redirects=True)
    body = response.get_data(as_text=True)
    assert "الحسابين" in body
    with flask_app.app_context():
        assert ProgressPayment.query.get(payment_id) is None
        leftover = JournalEntry.query.filter(
            JournalEntry.description.like(f"%PP-AUTO:{payment_id}%")
        ).count()
        assert leftover == 0
        assert JournalEntry.query.filter(
            (JournalEntry.debit_account_id == debit_id) | (JournalEntry.credit_account_id == debit_id),
            JournalEntry.description.like("%بند%"),
        ).count() == 0
        assert JournalEntry.query.filter(
            (JournalEntry.debit_account_id == credit_id) | (JournalEntry.credit_account_id == credit_id),
            JournalEntry.description.like("%بند%"),
        ).count() == 0


def test_supplier_payment_edit_and_delete_syncs_both_accounts(client):
    with flask_app.app_context():
        supplier = Supplier(name="مورد سداد متزامن")
        db.session.add(supplier)
        db.session.commit()
        sync_journal_related_accounts()
        supplier_id = supplier.id
        treasury = ChartOfAccount.query.filter_by(code="TRS-MAIN").first()
        treasury_id = treasury.id
        party = ChartOfAccount.query.filter_by(code=f"SUP-{supplier.id:04d}").first()
        party_id = party.id

    create = client.post("/supplier_payments", data={
        "supplier_id": str(supplier_id),
        "date": "2026-09-05",
        "amount": "500",
        "payment_method": "نقدي",
        "treasury_account_id": str(treasury_id),
        "reference": "PAY-SYNC",
    }, follow_redirects=True)
    assert create.status_code == 200

    with flask_app.app_context():
        payment = SupplierPayment.query.filter_by(supplier_id=supplier_id).first()
        payment_id = payment.id
        journal = JournalEntry.query.filter(JournalEntry.description.like(f"%PAY-AUTO:{payment_id}%")).first()
        assert journal is not None
        assert journal.debit_account_id == party_id
        assert journal.credit_account_id == treasury_id
        assert round(journal.amount, 2) == 500

    client.post(f"/supplier_payments/{payment_id}/update", data={
        "supplier_id": str(supplier_id),
        "date": "2026-09-05",
        "amount": "350",
        "payment_method": "نقدي",
        "treasury_account_id": str(treasury_id),
        "reference": "PAY-SYNC",
    }, follow_redirects=True)

    with flask_app.app_context():
        payment = SupplierPayment.query.get(payment_id)
        journal = JournalEntry.query.filter(JournalEntry.description.like(f"%PAY-AUTO:{payment_id}%")).first()
        assert round(payment.amount, 2) == 350
        assert journal is not None
        assert round(journal.amount, 2) == 350
        assert journal.debit_account_id == party_id
        assert journal.credit_account_id == treasury_id

    party_stmt = client.get(f"/accounts/{party_id}/statement").get_data(as_text=True)
    treasury_stmt = client.get(f"/accounts/{treasury_id}/statement").get_data(as_text=True)
    assert "350" in party_stmt.replace(",", "")
    assert "350" in treasury_stmt.replace(",", "")

    client.post(f"/documents/supplier_payment/{payment_id}/delete", follow_redirects=True)
    with flask_app.app_context():
        assert SupplierPayment.query.get(payment_id) is None
        assert JournalEntry.query.filter(JournalEntry.description.like(f"%PAY-AUTO:{payment_id}%")).count() == 0

    party_stmt = client.get(f"/accounts/{party_id}/statement").get_data(as_text=True)
    treasury_stmt = client.get(f"/accounts/{treasury_id}/statement").get_data(as_text=True)
    assert "PAY-SYNC" not in party_stmt
    assert "PAY-SYNC" not in treasury_stmt


def test_cannot_delete_account_or_posted_journal(client):
    with flask_app.app_context():
        treasury_id, expense_id = _ids()

    client.post("/journal", data={
        "date": "2026-09-05",
        "entry_action": "post",
        "line_description": ["قيد يُحذف من الحسابين"],
        "line_debit_account_id": [str(expense_id)],
        "line_credit_account_id": [str(treasury_id)],
        "line_amount": ["40"],
    }, follow_redirects=True)

    with flask_app.app_context():
        entry = JournalEntry.query.filter(JournalEntry.description == "قيد يُحذف من الحسابين").first()
        entry_id = entry.id

    delete_account = client.post(f"/accounts/{expense_id}/delete", follow_redirects=True)
    assert "لا يمكن حذف الحسابات" in delete_account.get_data(as_text=True)

    create_unused = client.post("/accounts", data={
        "code": "TST-DEL",
        "name": "حساب للتجربة",
        "category": "المصروفات",
        "opening_balance": "0",
    }, follow_redirects=True)
    assert create_unused.status_code == 200
    with flask_app.app_context():
        unused = ChartOfAccount.query.filter_by(code="TST-DEL").first()
        unused_id = unused.id
    deleted_unused = client.post(f"/accounts/{unused_id}/delete", follow_redirects=True)
    assert "تم حذف الحساب TST-DEL" in deleted_unused.get_data(as_text=True)
    with flask_app.app_context():
        assert ChartOfAccount.query.filter_by(code="TST-DEL").first() is None

    delete_journal = client.post(f"/journal/{entry_id}/delete", follow_redirects=True)
    assert "الحساب" in delete_journal.get_data(as_text=True)
    with flask_app.app_context():
        assert ChartOfAccount.query.get(expense_id) is not None
        assert JournalEntry.query.get(entry_id) is None


def test_client_name_spelling_reuses_same_account(client):
    with flask_app.app_context():
        from services.accounting import get_or_create_client_account
        first = get_or_create_client_account("شركة النور")
        db.session.commit()
        first_id = first.id
        second = get_or_create_client_account("  شركة   النور  ")
        db.session.commit()
        assert second.id == first_id
        matches = [
            item for item in ChartOfAccount.query.filter_by(category="العملاء").all()
            if "النور" in (item.name or "")
        ]
        assert len(matches) == 1


def test_change_password_page_available(client):
    response = client.get("/account/password")
    body = response.get_data(as_text=True)
    assert response.status_code == 200
    assert "حسابي" in body
    assert "تغيير كلمة السر" in body
    assert "تغيير باسورد الـ PIN" in body
    assert "123456" in body


def test_manual_journal_to_subcontractor_creates_advance(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-J1", project_name="ربط قيود", client_name="عميل ز",
            contract_type="مقاولات عامة", contract_value=500000,
        )
        sub = Subcontractor(name="مقاول يومية")
        db.session.add_all([project, sub])
        db.session.commit()
        sync_journal_related_accounts()
        sub_id = sub.id
        project_id = project.id
        sub_account = ChartOfAccount.query.filter_by(code=f"SUB-{sub.id:04d}").first()
        treasury = ChartOfAccount.query.filter_by(code="TRS-MAIN").first()
        sub_account_id = sub_account.id
        treasury_id = treasury.id

    client.post("/journal", data={
        "date": "2026-08-12",
        "entry_action": "post",
        "line_description": ["صرف لمقاول يومية"],
        "line_debit_account_id": [str(sub_account_id)],
        "line_credit_account_id": [str(treasury_id)],
        "line_amount": ["2500"],
    }, follow_redirects=True)

    with flask_app.app_context():
        advance = SubcontractorPayment.query.filter_by(subcontractor_id=sub_id).first()
        assert advance is not None
        assert round(advance.amount, 2) == 2500.0
        assert "JRN-SRC:" in (advance.notes or "")
        assert get_subcontractor_outstanding_advances(sub_id) == 2500.0
        statement = build_subcontractor_statement(Subcontractor.query.get(sub_id))
        assert statement["total_paid"] == 2500.0
        auto_dupes = JournalEntry.query.filter(JournalEntry.description.like("%SPAY-AUTO:%")).count()
        assert auto_dupes == 0

    client.post("/progress_payments", data={
        "project_id": str(project_id),
        "subcontractor_id": str(sub_id),
        "date": "2026-08-13",
        "retention_percentage": "0",
        "tax_percentage": "0",
        "advance_deduction": "2500",
        "description": ["أعمال"],
        "unit": ["متر مكعب"],
        "quantity": ["10"],
        "unit_price": ["400"],
    }, follow_redirects=True)

    with flask_app.app_context():
        payment = ProgressPayment.query.filter_by(subcontractor_id=sub_id).first()
        assert payment is not None
        assert round(payment.advance_deduction, 2) == 2500.0
        assert round(payment.net_value, 2) == 1500.0
        statement = build_subcontractor_statement(Subcontractor.query.get(sub_id))
        assert statement["total_works"] == 4000.0
        assert statement["total_paid"] == 2500.0
        assert statement["net_due"] == 1500.0


def test_manual_journal_to_supplier_creates_payment(client):
    with flask_app.app_context():
        supplier = Supplier(name="مورد يومية")
        db.session.add(supplier)
        db.session.commit()
        sync_journal_related_accounts()
        supplier_id = supplier.id
        sup_account = ChartOfAccount.query.filter_by(code=f"SUP-{supplier.id:04d}").first()
        treasury = ChartOfAccount.query.filter_by(code="TRS-MAIN").first()
        sup_account_id = sup_account.id
        treasury_id = treasury.id

    client.post("/journal", data={
        "date": "2026-08-12",
        "entry_action": "post",
        "line_description": ["سداد مورد"],
        "line_debit_account_id": [str(sup_account_id)],
        "line_credit_account_id": [str(treasury_id)],
        "line_amount": ["800"],
    }, follow_redirects=True)

    with flask_app.app_context():
        payment = SupplierPayment.query.filter_by(supplier_id=supplier_id).first()
        assert payment is not None
        assert round(payment.amount, 2) == 800.0
        assert JournalEntry.query.filter(JournalEntry.description.like("%PAY-AUTO:%")).count() == 0


def test_account_statement_lists_posted_moves(client):
    with flask_app.app_context():
        treasury_id, expense_id = _ids()
        treasury_name = ChartOfAccount.query.get(treasury_id).name

    listing = client.get("/accounts")
    assert listing.status_code == 200
    assert "كشف حساب" in listing.get_data(as_text=True)

    client.post("/journal", data={
        "date": "2026-08-20",
        "entry_action": "post",
        "line_description": ["صرف من الخزنة"],
        "line_debit_account_id": [str(expense_id)],
        "line_credit_account_id": [str(treasury_id)],
        "line_amount": ["250"],
    }, follow_redirects=True)

    response = client.get(f"/accounts/{treasury_id}/statement")
    body = response.get_data(as_text=True)
    assert response.status_code == 200
    assert "كشف حساب" in body
    assert treasury_name in body
    assert "250" in body.replace(",", "").replace(".", "")
    assert "صرف من الخزنة" in body


def test_journal_stamps_actor_and_writes_activity_log(client):
    listing = client.get("/journal")
    assert listing.status_code == 200
    assert "المستخدم" in listing.get_data(as_text=True)

    with flask_app.app_context():
        treasury_id, expense_id = _ids()

    client.post("/journal", data={
        "date": "2026-08-21",
        "entry_action": "post",
        "line_description": ["صرف تجريبي للتتبع"],
        "line_debit_account_id": [str(expense_id)],
        "line_credit_account_id": [str(treasury_id)],
        "line_amount": ["75"],
    }, follow_redirects=True)

    with flask_app.app_context():
        entry = JournalEntry.query.order_by(JournalEntry.id.desc()).first()
        assert entry is not None
        assert entry.created_by_name == "مدير النظام"
        assert entry.created_at
        log = ActivityLog.query.filter_by(entity_type="JournalEntry", entity_id=entry.id).first()
        assert log is not None
        assert log.user_name == "مدير النظام"
        assert log.action == "إضافة"
        assert log.summary

    page = client.get("/activity-log")
    body = page.get_data(as_text=True)
    assert page.status_code == 200
    assert "سجل حركات المستخدمين" in body
    assert "مدير النظام" in body
    assert "قيد يومية" in body


def test_data_entry_cannot_open_activity_log(client):
    with flask_app.app_context():
        clerk = User(username="clerk2", full_name="مدخل بيانات ٢", role=ROLE_DATA_ENTRY, is_active=True)
        clerk.set_password("secret12")
        db.session.add(clerk)
        db.session.commit()
        clerk_id = clerk.id

    probe = flask_app.test_client()
    with probe.session_transaction() as sess:
        sess["user_id"] = clerk_id
    response = probe.get("/activity-log", follow_redirects=True)
    body = response.get_data(as_text=True)
    assert "سجل حركات المستخدمين" not in body
    assert "ليست لديك صلاحية" in body or "متاح لمدير النظام فقط" in body


def test_account_rename_does_not_change_posted_amount(client):
    with flask_app.app_context():
        treasury_id, expense_id = _ids()
        expense = ChartOfAccount.query.get(expense_id)
        original_amount_account_id = expense.id
        original_opening = expense.opening_balance or 0

    client.post("/journal", data={
        "date": "2026-09-05",
        "entry_action": "post",
        "line_description": ["صرف ثابت"],
        "line_debit_account_id": [str(expense_id)],
        "line_credit_account_id": [str(treasury_id)],
        "line_amount": ["120"],
    }, follow_redirects=True)

    client.post(f"/accounts/{expense_id}/update", data={
        "code": "EXP-RENAMED",
        "name": "مصروف معاد تسميته",
        "category": "المصروفات",
        "opening_balance": str(original_opening),
    }, follow_redirects=True)

    with flask_app.app_context():
        entry = JournalEntry.query.filter(JournalEntry.description == "صرف ثابت").first()
        assert entry is not None
        assert round(entry.amount, 2) == 120
        assert entry.debit_account_id == original_amount_account_id
        account = ChartOfAccount.query.get(expense_id)
        assert account.name == "مصروف معاد تسميته"
        assert account.code == "EXP-RENAMED"
        assert account.category == "المصروفات"


def test_journal_account_picker_has_word_search(client):
    page = client.get("/journal")
    body = page.get_data(as_text=True)
    script = client.get("/static/searchable-select.js").get_data(as_text=True)
    assert page.status_code == 200
    assert "js-searchable-select" in body
    assert "ابحث بالكلمة" not in script
    assert "select.js-searchable-select" in script
    assert "searchable-select-input" in script
    assert "searchable-select-menu" in script


def test_employee_display_name_includes_tag(client):
    response = client.post("/employees", data={
        "name": "علي مهران",
        "tag": "مورد",
        "basic_salary": "1000",
    }, follow_redirects=True)
    assert response.status_code == 200
    body = response.get_data(as_text=True)
    assert "علي مهران (مورد)" in body
    with flask_app.app_context():
        employee = Employee.query.filter_by(name="علي مهران").first()
        assert employee is not None
        assert employee.display_name == "علي مهران (مورد)"
        account = ChartOfAccount.query.filter_by(code=f"EMP-{employee.id:04d}").first()
        assert account is not None
        assert account.category == "الموظفين"
        assert "علي مهران (مورد)" in account.name


def test_payroll_accrual_stays_on_expense_when_cash_is_paid(client):
    client.post("/employees", data={
        "name": "علي مهران",
        "tag": "مورد",
        "basic_salary": "1000",
    }, follow_redirects=True)

    with flask_app.app_context():
        employee = Employee.query.filter_by(name="علي مهران").first()
        employee_id = employee.id
        treasury_id = ChartOfAccount.query.filter_by(code="TRS-MAIN").first().id

    client.post("/hr", data={
        "action": "post_payroll",
        "period_month": "2026-09",
        "posting_date": "2026-09-30",
        "employee_id": [str(employee_id)],
        "work_days": ["30"],
        "vacation_days": ["0"],
        "basic_salary": ["1000"],
        "overtime": ["0"],
        "incentives": ["0"],
        "delay_deduction": ["0"],
        "permission_deduction": ["0"],
        "other_deductions": ["0"],
        "advances": ["0"],
        "payroll_notes": [""],
    }, follow_redirects=True)

    with flask_app.app_context():
        expense = get_account_by_code("EXP-SAL")
        employee = Employee.query.get(employee_id)
        emp_account = get_or_create_employee_account(employee)
        balances = build_account_balances()
        assert expense is not None
        assert round(balances.get(expense.id, 0.0), 2) == 1000
        assert round(balances.get(emp_account.id, 0.0), 2) == -1000
        slip = PayrollSlip.query.filter_by(employee_id=employee_id, period_month="2026-09").first()
        assert slip is not None
        assert round(slip.net_salary, 2) == 1000
        journal = JournalEntry.query.filter(JournalEntry.description.like(f"%PAYROLL-AUTO:{slip.id}%")).first()
        assert journal is not None
        assert journal.debit_account_id == expense.id
        assert journal.credit_account_id == emp_account.id

    client.post("/hr", data={
        "action": "pay_salary",
        "period_month": "2026-09",
        "employee_id": str(employee_id),
        "date": "2026-09-30",
        "amount": "500",
        "payment_method": "نقدي",
        "treasury_account_id": str(treasury_id),
    }, follow_redirects=True)

    with flask_app.app_context():
        expense = get_account_by_code("EXP-SAL")
        employee = Employee.query.get(employee_id)
        emp_account = get_or_create_employee_account(employee)
        treasury = ChartOfAccount.query.filter_by(code="TRS-MAIN").first()
        balances = build_account_balances()
        assert round(balances.get(expense.id, 0.0), 2) == 1000
        assert round(balances.get(emp_account.id, 0.0), 2) == -500
        assert round(balances.get(treasury.id, 0.0), 2) == -500
        payment = EmployeeSalaryPayment.query.filter_by(employee_id=employee_id).first()
        assert payment is not None
        pay_journal = JournalEntry.query.filter(JournalEntry.description.like(f"%ESAL-AUTO:{payment.id}%")).first()
        assert pay_journal is not None
        assert pay_journal.debit_account_id == emp_account.id
        assert pay_journal.credit_account_id == treasury.id

    page = client.get("/reports/profit_loss")
    body = page.get_data(as_text=True)
    assert "مصروف المرتبات" in body
    assert "1.000" in body


def test_excel_import_adds_supplier_and_employee(client):
    from io import BytesIO
    from openpyxl import Workbook
    from services.workbook import SHEET_EMPLOYEES, SHEET_SUPPLIERS

    workbook = Workbook()
    workbook.remove(workbook.active)
    suppliers = workbook.create_sheet(SHEET_SUPPLIERS)
    suppliers.append(["الاسم", "الكود", "التصنيف", "الاتصال", "ملاحظات"])
    suppliers.append(["مورد الإكسل", "", "مورد توريد مواد بناء", "", ""])
    employees = workbook.create_sheet(SHEET_EMPLOYEES)
    employees.append(["الاسم", "الميزة", "المرتب الأساسي", "الاتصال", "ملاحظات"])
    employees.append(["أحمد فني", "فني", 2500, "", ""])
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    page = client.post(
        "/import-workbook",
        data={"workbook": (buffer, "shoghl.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert page.status_code == 200
    assert "تم النقل" in page.get_data(as_text=True)
    with flask_app.app_context():
        assert Supplier.query.filter_by(name="مورد الإكسل").first() is not None
        employee = Employee.query.filter_by(name="أحمد فني").first()
        assert employee is not None
        assert employee.display_name == "أحمد فني (فني)"
        assert ChartOfAccount.query.filter_by(code=f"EMP-{employee.id:04d}").first() is not None
        assert ChartOfAccount.query.filter_by(code=f"SUP-{Supplier.query.filter_by(name='مورد الإكسل').first().id:04d}").first() is not None


def test_desktop_shortcut_uses_current_host(client):
    page = client.get("/desktop-shortcut")
    assert page.status_code == 200
    assert "اختصار لاب إبراهيم" in page.get_data(as_text=True)
    file_resp = client.get("/desktop-shortcut/url")
    assert file_resp.status_code == 200
    body = file_resp.get_data(as_text=True)
    assert "[InternetShortcut]" in body
    assert "URL=" in body


def test_home_hides_laptop_shortcut_and_uses_account_picker(client):
    home = client.get("/")
    body = home.get_data(as_text=True)
    assert "اختصار لاب إبراهيم" not in body
    assert "اختصار اللاب" not in body
    assert "انقل شغلك من الإكسل" not in body
    assert "النسخ الاحتياطي" in body
    assert 'name="account_id"' in body
    assert "js-searchable-select" in body
    with flask_app.app_context():
        account = ChartOfAccount.query.filter_by(code="TRS-MAIN").first()
        assert account is not None
        account_id = account.id
    lookup = client.get(f"/accounts/lookup?account_id={account_id}")
    assert lookup.status_code == 302
    assert f"/accounts/{account_id}/statement" in lookup.headers["Location"]


def test_zero_advance_deduction_is_respected(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-ADV0", project_name="سلفة صفر", client_name="عميل س",
            contract_type="مقاولات عامة", contract_value=1,
        )
        sub = Subcontractor(name="مقاول سلفة")
        db.session.add_all([project, sub])
        db.session.commit()
        pid, sid = project.id, sub.id
        treasury_id, _ = _ids()

    client.post("/subcontractor_payments", data={
        "subcontractor_id": str(sid),
        "project_id": str(pid),
        "amount": "2000",
        "payment_kind": "تحت الحساب",
        "payment_method": "نقدي",
        "treasury_account_id": str(treasury_id),
        "date": "2026-08-01",
    }, follow_redirects=True)

    client.post("/progress_payments", data={
        "project_id": str(pid),
        "subcontractor_id": str(sid),
        "date": "2026-08-02",
        "retention_percentage": "0",
        "tax_percentage": "0",
        "advance_deduction": "0",
        "description": ["أعمال"],
        "unit": ["متر مكعب"],
        "quantity": ["10"],
        "unit_price": ["300"],
    }, follow_redirects=True)

    with flask_app.app_context():
        payment = ProgressPayment.query.filter_by(subcontractor_id=sid).first()
        assert round(payment.advance_deduction, 2) == 0
        assert round(payment.net_value, 2) == 3000.0
        assert get_subcontractor_outstanding_advances(sid) == 2000.0


def test_settlement_does_not_count_as_outstanding_advance(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-SET", project_name="صرف صافي", client_name="عميل ص",
            contract_type="مقاولات عامة", contract_value=1,
        )
        sub = Subcontractor(name="مقاول صافي")
        db.session.add_all([project, sub])
        db.session.commit()
        pid, sid = project.id, sub.id
        treasury_id, _ = _ids()

    client.post("/progress_payments", data={
        "project_id": str(pid),
        "subcontractor_id": str(sid),
        "date": "2026-08-01",
        "retention_percentage": "10",
        "tax_percentage": "0",
        "advance_deduction": "0",
        "description": ["أعمال"],
        "unit": ["متر مكعب"],
        "quantity": ["10"],
        "unit_price": ["1000"],
    }, follow_redirects=True)

    with flask_app.app_context():
        first = ProgressPayment.query.filter_by(subcontractor_id=sid).first()
        assert round(first.net_value, 2) == 9000.0
        first_id = first.id

    client.post("/subcontractor_payments", data={
        "subcontractor_id": str(sid),
        "project_id": str(pid),
        "progress_payment_id": str(first_id),
        "payment_kind": "صرف مستخلص",
        "amount": "9000",
        "payment_method": "نقدي",
        "treasury_account_id": str(treasury_id),
        "date": "2026-08-03",
        "redirect_to": "progress_payment_detail",
    }, follow_redirects=True)

    with flask_app.app_context():
        assert get_subcontractor_outstanding_advances(sid) == 0
        assert progress_net_remaining(ProgressPayment.query.get(first_id)) == 0

    client.post("/progress_payments", data={
        "project_id": str(pid),
        "subcontractor_id": str(sid),
        "date": "2026-08-10",
        "retention_percentage": "0",
        "tax_percentage": "0",
        "advance_deduction": "0",
        "description": ["أعمال تالية"],
        "unit": ["متر مكعب"],
        "quantity": ["5"],
        "unit_price": ["1000"],
    }, follow_redirects=True)

    with flask_app.app_context():
        second = ProgressPayment.query.filter_by(subcontractor_id=sid).order_by(ProgressPayment.id.desc()).first()
        assert round(second.advance_deduction, 2) == 0
        assert round(second.net_value, 2) == 5000.0
        statement = build_subcontractor_statement(Subcontractor.query.get(sid))
        assert statement["total_works"] == 15000.0
        assert statement["total_paid"] == 9000.0
        assert statement["net_due"] == 5000.0


def test_closed_period_does_not_save_progress_payment(client):
    with flask_app.app_context():
        period = AccountingPeriod(name="أغسطس", from_date="2026-08-01", to_date="2026-08-31", status="مغلقة")
        project = Project(
            code="PRJ-CL", project_name="فترة مغلقة", client_name="عميل ط",
            contract_type="مقاولات عامة", contract_value=1,
        )
        sub = Subcontractor(name="مقاول مغلق")
        db.session.add_all([period, project, sub])
        db.session.commit()
        pid, sid = project.id, sub.id

    response = client.post("/progress_payments", data={
        "project_id": str(pid),
        "subcontractor_id": str(sid),
        "date": "2026-08-15",
        "description": ["أعمال"],
        "unit": ["متر مكعب"],
        "quantity": ["1"],
        "unit_price": ["100"],
    }, follow_redirects=True)
    assert "مغلقة" in response.get_data(as_text=True) or "فترة" in response.get_data(as_text=True)
    with flask_app.app_context():
        assert ProgressPayment.query.filter_by(subcontractor_id=sid).count() == 0


def test_progress_number_does_not_reuse_existing(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-NUM", project_name="أرقام", client_name="عميل ع",
            contract_type="مقاولات عامة", contract_value=1,
        )
        sub = Subcontractor(name="مقاول أرقام")
        db.session.add_all([project, sub])
        db.session.commit()
        first = ProgressPayment(
            project_id=project.id, subcontractor_id=sub.id, payment_number="1", date="2026-08-01",
        )
        third = ProgressPayment(
            project_id=project.id, subcontractor_id=sub.id, payment_number="3", date="2026-08-02",
        )
        db.session.add_all([first, third])
        db.session.commit()
        assert generate_progress_payment_number(project.id, sub.id) == "4"


def test_retention_percent_zero_clears_manual_amount_when_submitted_zero(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-RET0", project_name="ضمان صفر", client_name="عميل غ",
            contract_type="مقاولات عامة", contract_value=1,
        )
        sub = Subcontractor(name="مقاول ضمان")
        db.session.add_all([project, sub])
        db.session.commit()
        pid, sid = project.id, sub.id

    client.post("/progress_payments", data={
        "project_id": str(pid),
        "subcontractor_id": str(sid),
        "date": "2026-08-01",
        "retention_percentage": "0",
        "discount_insurance": "0",
        "tax_percentage": "0",
        "description": ["أعمال"],
        "unit": ["متر مكعب"],
        "quantity": ["10"],
        "unit_price": ["100"],
    }, follow_redirects=True)

    with flask_app.app_context():
        payment = ProgressPayment.query.filter_by(subcontractor_id=sid).first()
        assert round(payment.discount_insurance or 0, 2) == 0


def test_retention_release_posts_and_reduces_held_amount(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-REL", project_name="إفراج", client_name="عميل ف",
            contract_type="مقاولات عامة", contract_value=1,
        )
        sub = Subcontractor(name="مقاول إفراج")
        db.session.add_all([project, sub])
        db.session.commit()
        pid, sid = project.id, sub.id

    client.post("/progress_payments", data={
        "project_id": str(pid),
        "subcontractor_id": str(sid),
        "date": "2026-08-01",
        "retention_percentage": "10",
        "description": ["أعمال"],
        "unit": ["متر مكعب"],
        "quantity": ["10"],
        "unit_price": ["1000"],
    }, follow_redirects=True)

    with flask_app.app_context():
        assert get_subcontractor_outstanding_retention(sid) == 1000.0

    client.post("/retention_releases", data={
        "subcontractor_id": str(sid),
        "project_id": str(pid),
        "amount": "400",
        "date": "2026-08-20",
    }, follow_redirects=True)

    with flask_app.app_context():
        assert get_subcontractor_outstanding_retention(sid) == 600.0
        assert SubcontractorRetentionRelease.query.filter_by(subcontractor_id=sid).count() == 1
        statement = build_subcontractor_statement(Subcontractor.query.get(sid))
        assert statement["total_retention_released"] == 400.0
        assert statement["retention_outstanding"] == 600.0
        assert JournalEntry.query.filter(JournalEntry.description.like("%RETREL-AUTO:%")).count() == 1


def test_previous_executed_quantity_excludes_current_certificate(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-QTY", project_name="كميات", client_name="عميل ك",
            contract_type="مقاولات عامة", contract_value=1,
        )
        sub = Subcontractor(name="مقاول كميات")
        db.session.add_all([project, sub])
        db.session.commit()
        pid, sid = project.id, sub.id

    client.post("/progress_payments", data={
        "project_id": str(pid),
        "subcontractor_id": str(sid),
        "date": "2026-08-01",
        "description": ["ميول خرسانية"],
        "unit": ["متر مسطح"],
        "quantity": ["20"],
        "unit_price": ["50"],
    }, follow_redirects=True)

    with flask_app.app_context():
        first = ProgressPayment.query.filter_by(subcontractor_id=sid).first()
        assert previous_executed_quantity(pid, sid, None, "ميول خرسانية", "متر مسطح", first.id) == 0
        assert previous_executed_quantity(pid, sid, None, "ميول خرسانية", "متر مسطح") == 20.0


def test_can_delete_user_employee_and_unused_parties(client):
    create_user = client.post("/users", data={
        "username": "tempuser",
        "full_name": "مستخدم مؤقت",
        "password": "secret12",
        "role": ROLE_DATA_ENTRY,
        "is_active": "on",
    }, follow_redirects=True)
    assert create_user.status_code == 200
    with flask_app.app_context():
        temp = User.query.filter_by(username="tempuser").first()
        temp_id = temp.id
        admin_id = User.query.filter_by(username="admin").first().id

    self_delete = client.post(f"/users/{admin_id}/delete", follow_redirects=True)
    assert "لا يمكن حذف حسابك" in self_delete.get_data(as_text=True)
    with flask_app.app_context():
        assert User.query.get(admin_id) is not None

    deleted_user = client.post(f"/users/{temp_id}/delete", follow_redirects=True)
    assert "تم حذف المستخدم tempuser" in deleted_user.get_data(as_text=True)
    with flask_app.app_context():
        assert User.query.filter_by(username="tempuser").first() is None

    client.post("/employees", data={
        "name": "موظف مؤقت",
        "tag": "موظف",
        "basic_salary": "0",
    }, follow_redirects=True)
    with flask_app.app_context():
        emp = Employee.query.filter_by(name="موظف مؤقت").first()
        emp_id = emp.id
        emp_account = ChartOfAccount.query.filter_by(code=f"EMP-{emp_id:04d}").first()
        emp_account_id = emp_account.id if emp_account else None

    deleted_emp = client.post(f"/employees/{emp_id}/delete", follow_redirects=True)
    assert "تم حذف الموظف" in deleted_emp.get_data(as_text=True)
    with flask_app.app_context():
        assert Employee.query.get(emp_id) is None
        if emp_account_id:
            assert ChartOfAccount.query.get(emp_account_id) is None

    client.post("/suppliers", data={"name": "مورد مؤقت", "contact_info": "", "notes": ""}, follow_redirects=True)
    with flask_app.app_context():
        supplier = Supplier.query.filter_by(name="مورد مؤقت").first()
        supplier_id = supplier.id
    deleted_supplier = client.post(f"/suppliers/{supplier_id}/delete", follow_redirects=True)
    assert "تم حذف المورد" in deleted_supplier.get_data(as_text=True)
    with flask_app.app_context():
        assert Supplier.query.get(supplier_id) is None

    client.post("/subcontractors", data={"name": "مقاول مؤقت", "entity_kind": "مقاول تنفيذي (مصنعية ومعدات/عمالة)"}, follow_redirects=True)
    with flask_app.app_context():
        sub = Subcontractor.query.filter_by(name="مقاول مؤقت").first()
        sub_id = sub.id
    deleted_sub = client.post(f"/subcontractors/{sub_id}/delete", follow_redirects=True)
    assert "تم حذف جهة التعامل" in deleted_sub.get_data(as_text=True)
    with flask_app.app_context():
        assert Subcontractor.query.get(sub_id) is None


def test_admin_assigns_custom_user_permissions(client):
    from services.authz import ALL, effective_perms

    with flask_app.app_context():
        clerk = User(username="permclerk", full_name="مدخل صلاحيات", role=ROLE_DATA_ENTRY, is_active=True)
        clerk.set_password("secret12")
        db.session.add(clerk)
        db.session.commit()
        clerk_id = clerk.id
        admin_id = User.query.filter_by(username="admin").first().id

    probe = flask_app.test_client()
    with probe.session_transaction() as sess:
        sess["user_id"] = clerk_id
    blocked = probe.get("/accounts", follow_redirects=True)
    assert "ليست لديك صلاحية" in blocked.get_data(as_text=True)

    granted = client.post(f"/users/{clerk_id}/permissions", data={
        "perm": ["accounts", "accounts.create", "accounts.update"],
    }, follow_redirects=True)
    assert "تم تحديث صلاحيات" in granted.get_data(as_text=True)

    listing = probe.get("/accounts")
    assert listing.status_code == 200
    assert "إضافة حساب" in listing.get_data(as_text=True)

    with flask_app.app_context():
        extra = ChartOfAccount(code="PERM-DEL", name="حساب صلاحيات", category="المصروفات", opening_balance=0)
        db.session.add(extra)
        db.session.commit()
        extra_id = extra.id
    probe.post(f"/accounts/{extra_id}/delete", follow_redirects=True)
    with flask_app.app_context():
        assert ChartOfAccount.query.get(extra_id) is not None

    locked = client.post(f"/users/{admin_id}/permissions", data={"perm": ["accounts"]}, follow_redirects=True)
    assert "لا يمكن تقييدها" in locked.get_data(as_text=True)
    with flask_app.app_context():
        admin = User.query.get(admin_id)
        assert ALL in effective_perms(admin)

    denied = probe.post(f"/users/{admin_id}/permissions", data={"perm": ["accounts"]}, follow_redirects=True)
    assert "مدير النظام فقط" in denied.get_data(as_text=True) or "ليست لديك صلاحية" in denied.get_data(as_text=True)


def test_open_purchase_order_excluded_from_project_cost(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-OPENPO", project_name="مفتوح", client_name="عميل مواد",
            contract_type="مقاولات عامة", contract_value=1,
        )
        supplier = Supplier(name="مورد أوامر مفتوحة")
        db.session.add_all([project, supplier])
        db.session.commit()
        pid, sid = project.id, supplier.id

    client.post("/purchase_orders", data={
        "project_id": str(pid),
        "supplier_id": str(sid),
        "item_name": "أسمنت",
        "warehouse_name": "رئيسي",
        "quantity": "10",
        "unit_price": "50",
        "discount": "0",
        "date": "2026-09-01",
        "status": "مفتوح",
    }, follow_redirects=True)

    with flask_app.app_context():
        breakdown = build_project_cost_breakdown(Project.query.get(pid))
        assert breakdown["material_cost"] == 0.0
        order = PurchaseOrder.query.filter_by(project_id=pid).first()
        oid = order.id

    client.post(f"/purchase_orders/{oid}/update", data={
        "project_id": str(pid),
        "supplier_id": str(sid),
        "item_name": "أسمنت",
        "warehouse_name": "رئيسي",
        "quantity": "10",
        "unit_price": "50",
        "discount": "0",
        "date": "2026-09-01",
        "status": "مغلق",
    }, follow_redirects=True)

    with flask_app.app_context():
        breakdown = build_project_cost_breakdown(Project.query.get(pid))
        assert breakdown["material_cost"] == 500.0


def test_journal_reverse_blocked_when_source_period_closed(client):
    with flask_app.app_context():
        treasury_id, expense_id = _ids()

    client.post("/journal", data={
        "date": "2026-08-15",
        "entry_action": "post",
        "line_description": ["قيد للعكس"],
        "line_debit_account_id": [str(expense_id)],
        "line_credit_account_id": [str(treasury_id)],
        "line_amount": ["100"],
    }, follow_redirects=True)

    with flask_app.app_context():
        entry = JournalEntry.query.first()
        assert entry is not None
        eid = entry.id
        db.session.add(AccountingPeriod(
            name="أغسطس عكس", from_date="2026-08-01", to_date="2026-08-31", status="مغلقة",
        ))
        db.session.commit()

    response = client.post(f"/journal/{eid}/reverse", follow_redirects=True)
    assert "مغلقة" in response.get_data(as_text=True)
    with flask_app.app_context():
        assert JournalEntry.query.filter(JournalEntry.description.like("%عكس قيد%")).count() == 0


def test_purchase_order_update_blocked_on_original_closed_date(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-POCLOSE", project_name="إقفال أمر", client_name="عميل إقفال",
            contract_type="مقاولات عامة", contract_value=1,
        )
        supplier = Supplier(name="مورد إقفال أمر")
        db.session.add_all([project, supplier])
        db.session.commit()
        pid, sid = project.id, supplier.id

    client.post("/purchase_orders", data={
        "project_id": str(pid),
        "supplier_id": str(sid),
        "item_name": "رمل",
        "warehouse_name": "رئيسي",
        "quantity": "2",
        "unit_price": "30",
        "discount": "0",
        "date": "2026-08-10",
        "status": "مغلق",
    }, follow_redirects=True)

    with flask_app.app_context():
        order = PurchaseOrder.query.filter_by(project_id=pid).first()
        oid = order.id
        original_qty = order.total_value
        db.session.add(AccountingPeriod(
            name="أغسطس أوامر", from_date="2026-08-01", to_date="2026-08-31", status="مغلقة",
        ))
        db.session.commit()

    response = client.post(f"/purchase_orders/{oid}/update", data={
        "project_id": str(pid),
        "supplier_id": str(sid),
        "item_name": "رمل",
        "warehouse_name": "رئيسي",
        "quantity": "9",
        "unit_price": "30",
        "discount": "0",
        "date": "2026-09-10",
        "status": "مغلق",
    }, follow_redirects=True)
    assert "مغلقة" in response.get_data(as_text=True)
    with flask_app.app_context():
        order = PurchaseOrder.query.get(oid)
        assert order.date == "2026-08-10"
        assert round(order.total_value, 2) == round(original_qty, 2)


def test_equipment_journal_keeps_original_date(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-EQP", project_name="معدات", client_name="عميل معدات",
            contract_type="مقاولات عامة", contract_value=1,
        )
        db.session.add(project)
        db.session.commit()
        pid = project.id

    client.post("/equipment", data={
        "name": "لودر اختبار",
        "purchase_cost": "10000",
        "operating_cost": "250",
        "maintenance": "40",
        "hours_used": "5",
        "project_id": str(pid),
    }, follow_redirects=True)

    with flask_app.app_context():
        item = Equipment.query.filter_by(name="لودر اختبار").first()
        assert item is not None
        journals = JournalEntry.query.filter(JournalEntry.description.like(f"%EQP-AUTO:{item.id}%")).all()
        assert journals
        for journal in journals:
            journal.date = "2026-01-20"
        db.session.commit()
        sync_equipment_journals(item)
        refreshed = JournalEntry.query.filter(JournalEntry.description.like(f"%EQP-AUTO:{item.id}%")).all()
        assert refreshed
        assert all(journal.date == "2026-01-20" for journal in refreshed)


def test_dashboard_outstanding_counts_subcontractor_certificates_only(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-OUT", project_name="مستحقات", client_name="عميل مستحقات",
            contract_type="مقاولات عامة", contract_value=1,
        )
        sub = Subcontractor(name="مقاول مستحقات")
        db.session.add_all([project, sub])
        db.session.commit()
        pid, sid = project.id, sub.id

    client.post("/progress_payments", data={
        "project_id": str(pid),
        "subcontractor_id": str(sid),
        "date": "2026-09-01",
        "description": ["أعمال"],
        "unit": ["مقطوعية"],
        "quantity": ["1"],
        "unit_price": ["800"],
    }, follow_redirects=True)
    client.post("/client_progress_payments", data={
        "project_id": str(pid),
        "date": "2026-09-01",
        "description": ["فاتورة عميل"],
        "unit": ["مقطوعية"],
        "quantity": ["1"],
        "unit_price": ["9000"],
    }, follow_redirects=True)

    with flask_app.app_context():
        sub_count = ProgressPayment.query.filter(
            ProgressPayment.subcontractor_id.isnot(None),
            ProgressPayment.net_value > 0,
        ).count()
        client_count = ProgressPayment.query.filter(
            ProgressPayment.subcontractor_id.is_(None),
            ProgressPayment.net_value > 0,
        ).count()
        assert sub_count == 1
        assert client_count == 1


def test_data_entry_cannot_delete_client_progress_payment(client):
    with flask_app.app_context():
        project = Project(
            code="PRJ-DELIPC", project_name="حذف مستخلص", client_name="عميل حذف",
            contract_type="مقاولات عامة", contract_value=1,
        )
        clerk = User(username="ipcclerk", full_name="مدخل مستخلص", role=ROLE_DATA_ENTRY, is_active=True)
        clerk.set_password("secret12")
        db.session.add_all([project, clerk])
        db.session.commit()
        pid, clerk_id = project.id, clerk.id

    client.post("/client_progress_payments", data={
        "project_id": str(pid),
        "date": "2026-09-01",
        "description": ["بند عميل"],
        "unit": ["مقطوعية"],
        "quantity": ["1"],
        "unit_price": ["1200"],
    }, follow_redirects=True)

    with flask_app.app_context():
        ipc = ProgressPayment.query.filter(
            ProgressPayment.project_id == pid,
            ProgressPayment.subcontractor_id.is_(None),
        ).first()
        assert ipc is not None
        ipc_id = ipc.id

    probe = flask_app.test_client()
    with probe.session_transaction() as sess:
        sess["user_id"] = clerk_id
    response = probe.post(f"/documents/progress_payment/{ipc_id}/delete", follow_redirects=True)
    assert "ليست لديك صلاحية" in response.get_data(as_text=True)
    with flask_app.app_context():
        assert ProgressPayment.query.get(ipc_id) is not None


def test_problem_screens_include_required_fields(client):
    accounts = client.get("/accounts")
    assert "expense_class" in accounts.get_data(as_text=True)
    receipts = client.get("/client_receipts")
    assert "editRecNotes" in receipts.get_data(as_text=True)
    labor = client.get("/labor")
    assert labor.status_code == 200
    assert "إدارة العمالة" in labor.get_data(as_text=True)
    equipment = client.get("/equipment")
    assert equipment.status_code == 200
    estimations = client.get("/estimations")
    body = estimations.get_data(as_text=True)
    assert "est-line-total" in body
    assert "ipc-line" in body
    inventory = client.get("/inventory")
    assert "invLineTotal" in inventory.get_data(as_text=True)
    hr = client.get("/hr")
    assert hr.status_code == 200
    hr_body = hr.get_data(as_text=True)
    assert 'name="attendance_from"' in hr_body
    assert "إجازة مرضية" in hr_body
    assert "حفظ مسودة" in hr_body
    assert "طباعة كشف الحضور" in hr_body
    payments = client.get("/supplier_payments")
    assert "js-searchable-select" in payments.get_data(as_text=True)
    projects = client.get("/projects")
    assert projects.status_code == 200
    journal = client.get("/journal")
    assert "journal-line-row" in journal.get_data(as_text=True)


def test_admin_can_hide_dashboard_from_user(client):
    with flask_app.app_context():
        clerk = User(username="nodash", full_name="بدون رئيسية", role=ROLE_DATA_ENTRY, is_active=True)
        clerk.set_password("secret12")
        db.session.add(clerk)
        db.session.commit()
        clerk_id = clerk.id

    granted = client.post(f"/users/{clerk_id}/permissions", data={
        "perm": ["journal", "journal.view", "journal.create", "hr", "hr.create"],
    }, follow_redirects=True)
    assert "تم تحديث صلاحيات" in granted.get_data(as_text=True)

    probe = flask_app.test_client()
    with probe.session_transaction() as sess:
        sess["user_id"] = clerk_id
    home = probe.get("/", follow_redirects=True)
    body = home.get_data(as_text=True)
    assert "القيود اليومية" in body
    assert "ليست لديك صلاحية" in body
    journal = probe.get("/journal")
    assert journal.status_code == 200
    users_page = client.get("/users")
    assert "الشاشة الرئيسية" in users_page.get_data(as_text=True)


def test_payroll_sheet_matches_payslip_shape_and_prints(client):
    client.post("/employees", data={
        "name": "محمود الزين",
        "tag": "موظف",
        "job_title": "مندوب",
        "hometown": "الجيزة",
        "site": "القاهرة",
        "basic_salary": "7500",
    }, follow_redirects=True)
    page = client.get("/hr")
    body = page.get_data(as_text=True)
    assert "الراتب الأساسى" in body
    assert "موصلات" in body
    assert "الصافى النقدى" in body
    assert "تصدير Excel" in body

    with flask_app.app_context():
        employee_id = Employee.query.filter_by(name="محمود الزين").first().id

    saved = client.post("/hr", data={
        "action": "save_payroll",
        "period_month": "2026-09",
        "posting_date": "2026-09-20",
        "employee_id": [str(employee_id)],
        "work_days": ["30"],
        "vacation_days": ["0"],
        "basic_salary": ["7500"],
        "daily_rate": ["250"],
        "earned_salary": ["7500"],
        "transport": ["0"],
        "overtime": ["0"],
        "incentives": ["0"],
        "delay_deduction": ["0"],
        "permission_deduction": ["0"],
        "other_deductions": ["0"],
        "advances": ["0"],
        "payroll_notes": [""],
    }, follow_redirects=True)
    assert saved.status_code == 200

    with flask_app.app_context():
        slip = PayrollSlip.query.filter_by(employee_id=employee_id, period_month="2026-09").first()
        assert slip is not None
        assert round(slip.net_salary, 2) == 7500
        slip_id = slip.id

    payslip = client.get(f"/hr/payroll/{slip_id}/print")
    payslip_body = payslip.get_data(as_text=True)
    assert payslip.status_code == 200
    assert "قسيمة الراتب الشهرى" in payslip_body
    assert "محمود الزين" in payslip_body
    assert "مندوب" in payslip_body
    assert "IBRAHIM ABD ALKREEM" in payslip_body
    assert "شركة السيد الشيخ للمقاولات العمومية" in payslip_body
    assert "المعصرة" not in payslip_body
    assert "شارع العشرين" not in payslip_body
    assert "30 / 30" in payslip_body

    sheet = client.get("/hr/payroll/print?period_month=2026-09")
    sheet_body = sheet.get_data(as_text=True)
    assert sheet.status_code == 200
    assert "كشف المرتبات" in sheet_body
    assert "محمود الزين" in sheet_body
    assert "شركة السيد الشيخ للمقاولات العمومية" in sheet_body
    assert "7.500" in sheet_body


def test_payroll_absent_days_reduce_salary_automatically(client):
    client.post("/employees", data={
        "name": "سامي حضور",
        "tag": "موظف",
        "basic_salary": "3000",
    }, follow_redirects=True)
    with flask_app.app_context():
        employee_id = Employee.query.filter_by(name="سامي حضور").first().id

    saved = client.post("/hr", data={
        "action": "post_payroll",
        "period_month": "2026-09",
        "posting_date": "2026-09-30",
        "employee_id": [str(employee_id)],
        "work_days": ["25"],
        "vacation_days": ["0"],
        "basic_salary": ["3000"],
        "daily_rate": ["0"],
        "earned_salary": ["3000"],
        "transport": ["0"],
        "overtime": ["0"],
        "incentives": ["0"],
        "delay_deduction": ["0"],
        "permission_deduction": ["0"],
        "other_deductions": ["0"],
        "advances": ["0"],
        "payroll_notes": [""],
    }, follow_redirects=True)
    assert saved.status_code == 200

    with flask_app.app_context():
        slip = PayrollSlip.query.filter_by(employee_id=employee_id, period_month="2026-09").first()
        assert slip is not None
        assert round(slip.work_days, 2) == 25
        assert round(slip.daily_rate, 2) == 100
        assert round(slip.earned_salary, 2) == 2500
        assert round(slip.net_salary, 2) == 2500
        expense = get_account_by_code("EXP-SAL")
        balances = build_account_balances()
        assert round(balances.get(expense.id, 0.0), 2) == 2500
        slip_id = slip.id

    payslip = client.get(f"/hr/payroll/{slip_id}/print")
    body = payslip.get_data(as_text=True)
    assert payslip.status_code == 200
    assert "25 / 30" in body
    assert "2.500" in body
    assert "شركة السيد الشيخ للمقاولات العمومية" in body
    assert "print-watermark" in body


def test_attendance_range_sheet_posts_into_payroll(client):
    client.post("/employees", data={
        "name": "حسين فترة",
        "tag": "موظف",
        "basic_salary": "3000",
    }, follow_redirects=True)
    with flask_app.app_context():
        employee_id = Employee.query.filter_by(name="حسين فترة").first().id

    page = client.get("/hr?period_month=2026-09")
    body = page.get_data(as_text=True)
    assert "من تاريخ" in body
    assert "إجازة مرضية" in body
    assert "ترحيل إلى كشف المرتبات" in body

    draft = client.post("/hr", data={
        "action": "save_attendance",
        "period_month": "2026-09",
        "attendance_from": "2026-09-01",
        "attendance_to": "2026-09-05",
        "att_employee_id": [str(employee_id)],
        "att_status": ["حضور"],
        "att_from": ["2026-09-01"],
        "att_to": ["2026-09-05"],
        "att_check_in": [""],
        "att_check_out": [""],
        "att_notes": [""],
    }, follow_redirects=True)
    assert draft.status_code == 200
    draft_body = draft.get_data(as_text=True)
    assert "5 حضور" in draft_body
    assert "مسودة" in draft_body
    assert "طباعة كشف الحضور" in draft_body
    with flask_app.app_context():
        rows = EmployeeAttendance.query.filter_by(employee_id=employee_id).all()
        assert len(rows) == 5
        assert all(row.record_status == "مسودة" for row in rows)
        assert PayrollSlip.query.filter_by(employee_id=employee_id, period_month="2026-09").first() is None

    printed = client.get("/hr/attendance/print?period_month=2026-09")
    print_body = printed.get_data(as_text=True)
    assert printed.status_code == 200
    assert "حسين فترة" in print_body
    assert "5" in print_body
    assert "print-watermark" in print_body

    exported = client.get("/export/attendance?period_month=2026-09")
    assert exported.status_code == 200
    assert "spreadsheetml" in (exported.mimetype or "")

    posted = client.post("/hr", data={
        "action": "post_attendance",
        "period_month": "2026-09",
        "attendance_from": "2026-09-01",
        "attendance_to": "2026-09-05",
        "att_employee_id": [str(employee_id)],
        "att_status": ["إجازة مرضية"],
        "att_from": ["2026-09-01"],
        "att_to": ["2026-09-05"],
        "att_check_in": [""],
        "att_check_out": [""],
        "att_notes": ["مرض"],
    }, follow_redirects=True)
    assert posted.status_code == 200
    with flask_app.app_context():
        rows = EmployeeAttendance.query.filter_by(employee_id=employee_id).all()
        assert len(rows) == 5
        assert all(row.status == "إجازة مرضية" for row in rows)
        assert all(row.record_status == "مرحل" for row in rows)
        slip = PayrollSlip.query.filter_by(employee_id=employee_id, period_month="2026-09").first()
        assert slip is not None
        assert round(slip.work_days, 2) == 5
        assert round(slip.vacation_days, 2) == 5
        assert round(slip.earned_salary, 2) == 500
    assert "المعصرة" not in body


def test_attendance_month_end_posts_existing_drafts(client):
    client.post("/employees", data={
        "name": "سامي مسودة",
        "tag": "موظف",
        "basic_salary": "3000",
    }, follow_redirects=True)
    with flask_app.app_context():
        employee_id = Employee.query.filter_by(name="سامي مسودة").first().id

    client.post("/hr", data={
        "action": "save_attendance",
        "period_month": "2026-09",
        "att_employee_id": [str(employee_id)],
        "att_status": ["حضور"],
        "att_from": ["2026-09-01"],
        "att_to": ["2026-09-03"],
        "att_check_in": [""],
        "att_check_out": [""],
        "att_notes": [""],
    }, follow_redirects=True)

    posted = client.post("/hr", data={
        "action": "post_attendance",
        "period_month": "2026-09",
        "att_employee_id": [str(employee_id)],
        "att_status": [""],
        "att_from": ["2026-09-01"],
        "att_to": ["2026-09-03"],
        "att_check_in": [""],
        "att_check_out": [""],
        "att_notes": [""],
    }, follow_redirects=True)
    assert posted.status_code == 200
    with flask_app.app_context():
        rows = EmployeeAttendance.query.filter_by(employee_id=employee_id).all()
        assert len(rows) == 3
        assert all(row.status == "حضور" for row in rows)
        assert all(row.record_status == "مرحل" for row in rows)
        slip = PayrollSlip.query.filter_by(employee_id=employee_id, period_month="2026-09").first()
        assert slip is not None
        assert round(slip.work_days, 2) == 3


def test_employee_code_autofills_and_rejects_duplicates(client):
    first = client.post("/employees", data={
        "name": "موظف كود أول",
        "tag": "موظف",
    }, follow_redirects=True)
    assert first.status_code == 200
    with flask_app.app_context():
        first_emp = Employee.query.filter_by(name="موظف كود أول").first()
        assert first_emp.code
        taken_code = first_emp.code

    duplicate = client.post("/employees", data={
        "name": "موظف كود مكرر",
        "tag": "موظف",
        "code": taken_code,
    }, follow_redirects=True)
    body = duplicate.get_data(as_text=True)
    assert "مسجّل على الموظف" in body
    with flask_app.app_context():
        assert Employee.query.filter_by(name="موظف كود مكرر").first() is None

    second = client.post("/employees", data={
        "name": "موظف كود ثاني",
        "tag": "موظف",
    }, follow_redirects=True)
    assert second.status_code == 200
    with flask_app.app_context():
        second_emp = Employee.query.filter_by(name="موظف كود ثاني").first()
        assert second_emp is not None
        assert second_emp.code
        assert second_emp.code != taken_code
        assert first_emp.code == "EMP-0001"
        assert second_emp.code == "EMP-0002"


def test_employee_codes_rebuild_and_clear_continues_sequence(client):
    client.post("/employees", data={"name": "زيد قديم", "tag": "موظف", "code": "ZZ-9"}, follow_redirects=True)
    client.post("/employees", data={"name": "عمر قديم", "tag": "موظف", "code": "AA-1"}, follow_redirects=True)
    with flask_app.app_context():
        from services.accounting import EMPLOYEE_CODE_RESET_MARK, rebuild_employee_codes_once
        ActivityLog.query.filter_by(entity_label=EMPLOYEE_CODE_RESET_MARK).delete()
        db.session.commit()
        rebuild_employee_codes_once()
        db.session.commit()
        rows = Employee.query.order_by(Employee.id).all()
        assert [row.code for row in rows] == ["EMP-0001", "EMP-0002"]
        first_id = rows[0].id

    cleared = client.post(f"/employees/{first_id}/update", data={
        "name": "زيد قديم",
        "tag": "موظف",
        "code": "",
        "basic_salary": "0",
        "is_active": "on",
    }, follow_redirects=True)
    assert cleared.status_code == 200
    with flask_app.app_context():
        first = Employee.query.get(first_id)
        assert first.code == "EMP-0003"


def test_account_code_autosequence_and_rejects_duplicates(client):
    page = client.get("/accounts")
    assert "تلقائي إن تُرك فارغًا" in page.get_data(as_text=True)
    first = client.post("/accounts", data={
        "name": "حساب تسلسلي",
        "category": "المصروفات",
        "opening_balance": "0",
    }, follow_redirects=True)
    assert first.status_code == 200
    second = client.post("/accounts", data={
        "name": "حساب تسلسلي ٢",
        "category": "المصروفات",
        "opening_balance": "0",
    }, follow_redirects=True)
    assert second.status_code == 200
    with flask_app.app_context():
        one = ChartOfAccount.query.filter_by(name="حساب تسلسلي").first()
        two = ChartOfAccount.query.filter_by(name="حساب تسلسلي ٢").first()
        assert one.code == "EXP-0001"
        assert two.code == "EXP-0002"
        first_id = one.id

    duplicate = client.post("/accounts", data={
        "name": "حساب مكرر كود",
        "category": "المصروفات",
        "code": "EXP-0001",
        "opening_balance": "0",
    }, follow_redirects=True)
    assert "مسجّل على الحساب" in duplicate.get_data(as_text=True)

    cleared = client.post(f"/accounts/{first_id}/update", data={
        "name": "حساب تسلسلي",
        "category": "المصروفات",
        "code": "",
        "opening_balance": "0",
    }, follow_redirects=True)
    assert cleared.status_code == 200
    with flask_app.app_context():
        one = ChartOfAccount.query.get(first_id)
        assert one.code == "EXP-0003"


def test_data_entry_screens_export_excel(client):
    export = client.get("/export/employees")
    assert export.status_code == 200
    assert "spreadsheetml" in (export.mimetype or "")
    payroll = client.get("/export/payroll?period_month=2026-09")
    assert payroll.status_code == 200
    journal = client.get("/export/journal")
    assert journal.status_code == 200
    hr_page = client.get("/hr")
    assert "/export/payroll" in hr_page.get_data(as_text=True)
    employees_page = client.get("/employees")
    assert "/export/employees" in employees_page.get_data(as_text=True)


def test_sales_trip_sheet_totals_feed_driver_account(client):
    page = client.get("/sales")
    body = page.get_data(as_text=True)
    assert page.status_code == 200
    assert "رقم البون" in body
    assert "صافي الكمية" in body
    assert "تصدير Excel" in body
    assert "طباعة" in body
    assert "ابحث واختر العميل" in body
    assert "ابحث واختر السائق" in body
    assert "name=\"period_label\"" not in body
    assert "line_period_label" not in body

    saved = client.post("/sales", data={
        "line_date": ["2026-09-15", "2026-09-16"],
        "line_voucher_number": ["58927", "58928"],
        "line_trip_type": ["سيارة", "قلاب"],
        "line_client_name": ["الكتيبة", "الكتيبة"],
        "line_notes": ["كسارة القاهرة - الكورنيش", "كسارة القاهرة - المعادي"],
        "line_distance_km": ["110", "40"],
        "line_driver_name": ["شعبان محمد", "كريم علي"],
        "line_tractor_number": ["1412/3673", ""],
        "line_cubage": ["64", "20"],
        "line_discount": ["6", "0"],
        "line_unit_price": ["185", "150"],
        "line_advances": ["0", "0"],
    }, follow_redirects=True)
    assert saved.status_code == 200
    saved_body = saved.get_data(as_text=True)
    assert "58927" in saved_body
    assert "58928" in saved_body
    assert "شعبان محمد" in saved_body
    assert "كريم علي" in saved_body
    assert "10.730" in saved_body

    with flask_app.app_context():
        trip = SalesTrip.query.filter_by(voucher_number="58927").first()
        assert trip is not None
        assert round(trip.net_quantity, 2) == 58
        assert round(trip.total_amount, 2) == 10730
        assert round(trip.remaining, 2) == 10730
        revenue = get_account_by_code("REV-TRP")
        driver_account = ChartOfAccount.query.filter_by(name="سائق - شعبان محمد", category="السواقين").first()
        balances = build_account_balances()
        assert revenue is not None
        assert round(balances.get(revenue.id, 0.0), 2) == -13730
        assert driver_account is not None
        journal = JournalEntry.query.filter(JournalEntry.description.like(f"%SALE-AUTO:{trip.id}%")).first()
        assert journal is not None
        assert journal.credit_account_id == revenue.id
        assert round(balances.get(journal.debit_account_id, 0.0), 2) == 13730
        driver_account_id = driver_account.id

    drivers = client.get("/driver_compensation")
    drivers_body = drivers.get_data(as_text=True)
    assert "شعبان محمد" in drivers_body
    assert "نقلات المبيعات" in drivers_body
    assert "58927" in drivers_body

    driver_statement = client.get(f"/accounts/{driver_account_id}/statement")
    driver_statement_body = driver_statement.get_data(as_text=True)
    assert driver_statement.status_code == 200
    assert "نقلات السائق" in driver_statement_body
    assert "58927" in driver_statement_body
    assert "كسارة القاهرة - الكورنيش" in driver_statement_body

    printed = client.get("/sales/print")
    printed_body = printed.get_data(as_text=True)
    assert printed.status_code == 200
    assert "شركة السيد الشيخ للمقاولات العمومية" in printed_body
    assert "رقم البون" in printed_body
    assert "10.730" in printed_body

    export = client.get("/export/sales")
    assert export.status_code == 200
    assert "spreadsheetml" in (export.mimetype or "")


def test_admin_can_hide_main_treasury_from_user(client):
    with flask_app.app_context():
        main = ChartOfAccount.query.filter_by(code="TRS-MAIN").first()
        site = ChartOfAccount.query.filter_by(code="TRS-SITE").first()
        expense = ChartOfAccount.query.filter_by(code="EXP-ELC").first() or ChartOfAccount.query.filter_by(category="المصروفات").first()
        main_id, site_id, expense_id = main.id, site.id, expense.id
        clerk = User(username="mustafa", full_name="مصطفى", role=ROLE_DATA_ENTRY, is_active=True)
        clerk.set_password("secret12")
        db.session.add(clerk)
        db.session.commit()
        clerk_id = clerk.id

    users_page = client.get("/users")
    users_body = users_page.get_data(as_text=True)
    assert users_page.status_code == 200
    assert "الخزن الظاهرة للمستخدم" in users_body
    assert "TRS-MAIN" in users_body
    assert "modal-dialog-scrollable" in users_body

    saved = client.post(f"/users/{clerk_id}/permissions", data={
        "perm": ["dashboard", "accounts", "accounts.create", "journal", "journal.view", "journal.create"],
        "treasury_scope": "1",
        "visible_treasury": [str(site_id)],
    }, follow_redirects=True)
    assert "تم تحديث صلاحيات" in saved.get_data(as_text=True)

    client.post("/journal", data={
        "date": "2026-09-20",
        "entry_action": "post",
        "line_description": ["صرف من الخزنة الرئيسية للشيخ ابراهيم"],
        "line_reference": ["TRS-HIDE"],
        "line_debit_account_id": [str(expense_id)],
        "line_credit_account_id": [str(main_id)],
        "line_amount": ["2500"],
    }, follow_redirects=True)

    probe = flask_app.test_client()
    with probe.session_transaction() as sess:
        sess["user_id"] = clerk_id

    accounts_body = probe.get("/accounts").get_data(as_text=True)
    assert "TRS-MAIN" not in accounts_body
    assert "الخزنة الرئيسية" not in accounts_body

    blocked = probe.get(f"/accounts/{main_id}/statement", follow_redirects=True)
    assert "ليست لديك صلاحية" in blocked.get_data(as_text=True)

    journal_body = probe.get("/journal").get_data(as_text=True)
    assert "صرف من الخزنة الرئيسية للشيخ ابراهيم" not in journal_body
    assert "TRS-HIDE" not in journal_body

    rejected = probe.post("/journal", data={
        "date": "2026-09-21",
        "entry_action": "draft",
        "line_description": ["محاولة استخدام خزنة محجوبة"],
        "line_debit_account_id": [str(expense_id)],
        "line_credit_account_id": [str(main_id)],
        "line_amount": ["100"],
    }, follow_redirects=True)
    assert "ليست لديك صلاحية استخدام هذه الخزنة" in rejected.get_data(as_text=True)


def test_admin_amount_eye_needs_pin(client):
    locked_home = client.get("/")
    locked_body = locked_home.get_data(as_text=True)
    assert locked_home.status_code == 200
    assert "••••" in locked_body
    assert "js-secret-amount" in locked_body
    assert "wallet-eye-btn" in locked_body
    assert "walletPinModal" in locked_body

    accounts_page = client.get("/accounts")
    accounts_body = accounts_page.get_data(as_text=True)
    assert "wallet-eye-btn" in accounts_body
    assert "TRS-MAIN" in accounts_body

    wrong = client.post("/account/pin/unlock", data={"pin": "غلط"})
    assert wrong.status_code == 403

    unlocked = client.post("/account/pin/unlock", data={"pin": "123456"})
    assert unlocked.status_code == 200
    assert unlocked.get_json()["ok"] is True
    still_masked = client.get("/").get_data(as_text=True)
    assert "••••" in still_masked
    assert "js-secret-amount" in still_masked

    changed = client.post("/account/pin", data={
        "current_pin": "123456",
        "new_pin": "778899",
        "confirm_pin": "778899",
    }, follow_redirects=True)
    assert "تم حفظ الرقم السري" in changed.get_data(as_text=True)
    assert client.post("/account/pin/unlock", data={"pin": "123456"}).status_code == 403
    assert client.post("/account/pin/unlock", data={"pin": "778899"}).status_code == 200


def test_regular_user_has_own_treasury_pin(client):
    with flask_app.app_context():
        clerk = User(username="pinclerk", full_name="مدخل عين", role=ROLE_DATA_ENTRY, is_active=True)
        clerk.set_password("secret12")
        db.session.add(clerk)
        db.session.commit()
        clerk_id = clerk.id

    probe = flask_app.test_client()
    with probe.session_transaction() as sess:
        sess["user_id"] = clerk_id

    home = probe.get("/")
    home_body = home.get_data(as_text=True)
    assert home.status_code == 200
    assert "wallet-eye-btn" in home_body
    assert "حسابي" in home_body
    assert probe.post("/account/pin/unlock", data={"pin": "123456"}).status_code == 200
    assert probe.post("/account/pin/unlock", data={"pin": "ابراهيم"}).status_code == 403

    account_page = probe.get("/account/password")
    assert "تغيير باسورد الـ PIN" in account_page.get_data(as_text=True)
    saved = probe.post("/account/pin", data={
        "current_pin": "123456",
        "new_pin": "445566",
        "confirm_pin": "445566",
    }, follow_redirects=True)
    assert "تم حفظ الرقم السري" in saved.get_data(as_text=True)
    assert probe.post("/account/pin/unlock", data={"pin": "123456"}).status_code == 403
    assert probe.post("/account/pin/unlock", data={"pin": "445566"}).status_code == 200
    assert client.post("/account/pin/unlock", data={"pin": "123456"}).status_code == 200







